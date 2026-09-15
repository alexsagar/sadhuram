import docx
from pathlib import Path

doc = docx.Document("Details.docx")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")

for t_idx, table in enumerate(doc.tables):
    print(f"\n=== Table {t_idx}: {len(table.rows)} rows ===")
    for r_idx, row in enumerate(table.rows):
        cells = [c.text.replace("\n", " ").strip() for c in row.cells]
        # remove duplicate consecutive cells caused by merged cells
        unique_cells = []
        for c in cells:
            if not unique_cells or c != unique_cells[-1]:
                unique_cells.append(c)
        print(f"  [{r_idx}]: {' | '.join(unique_cells[:5])}")

# Check embedded images in the docx
import zipfile
with zipfile.ZipFile("Details.docx", "r") as z:
    media_files = [f for f in z.namelist() if f.startswith("word/media/")]
    print(f"\nEmbedded images in docx: {len(media_files)}")
    for mf in media_files:
        info = z.getinfo(mf)
        print(f"  {mf} ({info.file_size} bytes)")
